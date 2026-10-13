#!/usr/bin/env python3
"""Build the DIBY question-bank DOCX in the requested layout:

    Chapter, Lecture NN                 <- one title per chapter+lecture
    DIBY-02   ·  page 37                <- tag (+ where it comes from)
    <italic question with real typeset equations>          [Ans.: 13]
      options (small, grey)
    (blank line)  ... next DIBY ... then the next title

Equations are converted from the OCR text to LaTeX and rendered as native Word
equations (OMML) through pandoc, so they look typeset instead of "(x^m + 1/x^2)^22".
Section markers (slides with a DIBY mark but no question) are not included.
"""
from __future__ import annotations

import copy
import json
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from ocr_math_to_latex import segments          # noqa: E402

import pypandoc                                  # noqa: E402
from docx import Document                        # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT  # noqa: E402
from docx.oxml.ns import qn                      # noqa: E402
from docx.shared import Cm, Pt, RGBColor         # noqa: E402
from lxml import etree                           # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "deliverables/DIBY_Master.json")
OUT = os.path.join(ROOT, "deliverables/DIBY_Questions.docx")
CSV_OUT = os.path.join(ROOT, "deliverables/DIBY_Questions.csv")

CHAPTER_ORDER = ["Binomial Theorem", "Permutations & Combinations", "Quadratic Equations",
                 "Trigonometric Equation", "Trigonometric Functions"]
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

GREY = RGBColor(0x77, 0x77, 0x77)
ANSWER_BLUE = RGBColor(0x11, 0x55, 0x99)


# --------------------------------------------------------------------- helpers
def norm_tag(label: str, internal: str, marker: str) -> str:
    """Display tag: normalise '#DIBY-36' -> 'DIBY-36'."""
    lab = (label or "").strip().lstrip("#").strip()
    lab = re.sub(r"\s{2,}", " ", lab)
    if re.match(r"(?i)^diby\s*-\s*", lab):
        lab = re.sub(r"(?i)^diby\s*-\s*", "DIBY-", lab)
        lab = re.sub(r"(?i)^DIBY-(\d)\s*\(", r"DIBY-\1 (", lab)
    if not lab or lab.lower().startswith("none"):
        lab = internal.strip() or (marker.strip() if marker.strip() else "DIBY (no tag printed)")
    return lab


def clean_answer(ans: str) -> tuple[str, str]:
    """Return (short_inline_answer, long_note)."""
    ans = (ans or "").strip()
    if not ans:
        return "", ""
    if ans.lower().startswith("none"):
        return "", ""
    m = re.search(r"\[\s*Ans\.?\s*:?\s*([^\]]{1,40})\]", ans)
    if m:
        return f"[Ans.: {m.group(1).strip()}]", ""
    m = re.match(r"^Ans\.?\s*[:.]?\s*([^(\[]+?)\s*[(\[]", ans)
    if m and m.group(1).strip():
        return f"[Ans.: {m.group(1).strip()}]", ans
    m = re.match(r"^Ans\.?\s*[:.]?\s*([^(\[]+)$", ans)
    if m and m.group(1).strip():
        return f"[Ans.: {m.group(1).strip()}]", ""
    if len(ans) <= 24:
        return f"[Ans.: {ans}]" if not ans.startswith("[") else ans, ""
    return "", ans


def clean_options(opts: str) -> str:
    o = (opts or "").strip()
    if not o or o.lower().startswith("none") or "no options" in o.lower():
        return ""
    if o.lower().startswith("not extracted"):
        return ""
    return o


# ------------------------------------------------------- batch LaTeX -> OMML
def latex_to_omml(latex_list: list[str]) -> dict[str, etree._Element]:
    """Convert many LaTeX strings to OMML in ONE pandoc run."""
    uniq = sorted(set(latex_list))
    if not uniq:
        return {}
    lines = []
    for i, tex in enumerate(uniq):
        lines.append(f"ZQ{i:05d} ${tex}$")
        lines.append("")
    with tempfile.TemporaryDirectory() as td:
        md = os.path.join(td, "m.md")
        dx = os.path.join(td, "m.docx")
        open(md, "w").write("\n".join(lines))
        pypandoc.convert_file(md, "docx", outputfile=dx,
                              extra_args=["--standalone", "--quiet"])
        with zipfile.ZipFile(dx) as z:
            root = etree.fromstring(z.read("word/document.xml"))
    out: dict[str, etree._Element] = {}
    for p in root.iter(f"{{{W_NS}}}p"):
        text = "".join(t.text or "" for t in p.iter(f"{{{W_NS}}}t"))
        m = re.search(r"ZQ(\d{5})", text)
        if not m:
            continue
        omml = p.find(f"{{{M_NS}}}oMath")
        if omml is None:
            continue
        out[uniq[int(m.group(1))]] = omml
    return out


# ------------------------------------------------------------------ docx bits
def add_math(par, omml):
    par._p.append(copy.deepcopy(omml))


def styled_heading(doc, text, size=13, space_before=16):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.italic = True
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    return p


def main():
    recs = json.load(open(SRC))
    items = [r for r in recs if r["prompt"]]
    items.sort(key=lambda r: (CHAPTER_ORDER.index(r["chapter"]) if r["chapter"] in CHAPTER_ORDER else 99,
                              r["lecture"], r["pdf_page"], r["master_id"]))

    # 1) work out math spans for question + options text
    prepared = []
    all_latex = []
    for r in items:
        q_segs = segments(r["prompt"])
        o_txt = clean_options(r["options"])
        o_segs = segments(o_txt) if o_txt else []
        prepared.append((r, q_segs, o_segs))
        all_latex += [v for k, v in q_segs if k == "math"]
        all_latex += [v for k, v in o_segs if k == "math"]
    omml = latex_to_omml(all_latex)
    print(f"maths spans: {len(set(all_latex))} unique, {len(omml)} rendered as Word equations")

    # 2) build the document
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.7)
    text_width = Cm(17.0)

    n = doc.styles["Normal"]
    n.font.name = "Calibri"
    n.font.size = Pt(10.5)
    n.paragraph_format.space_after = Pt(0)

    # title block
    t = doc.add_paragraph()
    t.paragraph_format.space_before = Pt(2)
    tr = t.add_run("DIBY — Question Bank")
    tr.bold = True
    tr.italic = True
    tr.font.size = Pt(17)
    tr.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(10)
    qc = Counter(r["prompt_source"].split("(")[0].strip() for r, _, _ in prepared)
    n_ver = qc.get("verified visual transcription", 0)
    n_ext = qc.get("dual-pass OCR extraction, not visually QA'd", 0)
    n_drf = qc.get("raw OCR reconstructed from the same page", 0)
    sr = sub.add_run(f"{len(items)} questions · chapters in order · "
                     f"{n_ver} verified, {n_ext} ocr-extract, {n_drf} ocr-draft (unverified)")
    sr.italic = True
    sr.font.size = Pt(9)
    sr.font.color.rgb = GREY

    prev_key = None
    for r, q_segs, o_segs in prepared:
        key = (r["chapter"], r["lecture"])
        if key != prev_key:
            styled_heading(doc, f"{r['chapter']}, Lecture {r['lecture']}")
            prev_key = key

        # tag line
        tag = doc.add_paragraph()
        tag.paragraph_format.space_before = Pt(8)
        tag.paragraph_format.space_after = Pt(2)
        tag.paragraph_format.keep_with_next = True
        tr = tag.add_run(norm_tag(r["printed_label"], r["internal_id"], r["marker_text"]))
        tr.bold = True
        tr.font.size = Pt(11)
        pg = tag.add_run(f"   ·   {r['source_pdf']} — PDF page {r['pdf_page']}")
        pg.font.size = Pt(8)
        pg.italic = True
        pg.font.color.rgb = GREY

        # question paragraph (italic, maths inline, answer right-aligned)
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(2)
        q.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
        for kind, val in q_segs:
            if kind == "math" and val in omml:
                add_math(q, omml[val])
            else:
                run = q.add_run(val)
                run.italic = True
        short, long_note = clean_answer(r["answer"])
        if short:
            tab = q.add_run("\t" + short)
            tab.italic = True
            tab.font.color.rgb = ANSWER_BLUE

        # options line
        if o_segs:
            op = doc.add_paragraph()
            op.paragraph_format.space_after = Pt(2)
            op.paragraph_format.left_indent = Cm(0.5)
            for kind, val in o_segs:
                if kind == "math" and val in omml:
                    add_math(op, omml[val])
                else:
                    run = op.add_run(val)
                    run.italic = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = GREY
            for run in op.runs:
                run.font.size = Pt(9.5)
                run.font.color.rgb = GREY

        # long answer note
        if long_note:
            an = doc.add_paragraph()
            an.paragraph_format.space_after = Pt(2)
            an.paragraph_format.left_indent = Cm(0.5)
            ar = an.add_run(long_note)
            ar.italic = True
            ar.font.size = Pt(9)
            ar.font.color.rgb = GREY

        doc.add_paragraph().paragraph_format.space_after = Pt(6)   # spacer

    doc.save(OUT)
    print("wrote", OUT, os.path.getsize(OUT), "bytes")

    # 3) machine-readable side-car of exactly what is in the document
    import csv
    with open(CSV_OUT, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["document_order", "tag", "chapter", "lecture", "source_pdf", "pdf_page",
                    "question_text", "options", "answer_shown", "text_quality", "master_id"])
        for i, (r, _, _) in enumerate(prepared, 1):
            w.writerow([i, norm_tag(r["printed_label"], r["internal_id"], r["marker_text"]),
                        r["chapter"], r["lecture"], r["source_pdf"], r["pdf_page"],
                        r["prompt"], clean_options(r["options"]), r["answer"],
                        r["prompt_source"], r["master_id"]])
    print("wrote", CSV_OUT)


if __name__ == "__main__":
    main()
