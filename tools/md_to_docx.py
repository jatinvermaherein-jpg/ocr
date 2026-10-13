"""Convert a Twin 1 markdown transcript to DOCX (plan allows DOCX or Markdown; we ship both)."""
import sys, re, os
from docx import Document
from docx.shared import Pt

def add_runs(par, text):
    # split on **bold** and *italic*
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok: continue
        if tok.startswith("**") and tok.endswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok.replace("`", ""))

def main():
    src, dst = sys.argv[1], sys.argv[2]
    doc = Document(); doc.styles["Normal"].font.size = Pt(10)
    for raw in open(src):
        line = raw.rstrip("\n")
        s = line.strip()
        if not s or s in ("---",) or s.startswith("<a id="):
            continue
        if s.startswith("# "):
            doc.add_heading(s[2:].replace("**",""), 0)
        elif s.startswith("## "):
            doc.add_heading(re.sub(r"[`*]", "", s[3:]), 1)
        elif s.startswith("### "):
            doc.add_heading(re.sub(r"[`*]", "", s[4:]), 2)
        elif s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, s[2:])
        elif re.match(r"^\|", s):
            cells = [re.sub(r"[`*]", "", c).strip() for c in s.strip("|").split("|")]
            if all(set(c) <= set("-: ") for c in cells):  # separator row
                continue
            tb = doc.tables[-1] if doc.tables and getattr(doc, "_last_table_rows", None) == "tbl" else None
            p = doc.add_paragraph(" | ".join(cells), style="No Spacing")
        else:
            p = doc.add_paragraph(); add_runs(p, s)
    doc.save(dst)
    print("wrote", dst, os.path.getsize(dst), "bytes")

if __name__ == "__main__":
    main()
