#!/usr/bin/env python3
"""Build the question-bank Word document: one entry per DIBY item showing
Question · Options · Answer · Tag · where it comes from.

Reads deliverables/DIBY_Master.json (built by tools/build_diby_master.py).
Writes deliverables/DIBY_Questions.docx
"""
from __future__ import annotations

import json
import os
from collections import Counter

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "deliverables/DIBY_Master.json")
OUT = os.path.join(ROOT, "deliverables/DIBY_Questions.docx")

CHAPTER_ORDER = ["Binomial Theorem", "Permutations & Combinations", "Quadratic Equations",
                 "Trigonometric Equation", "Trigonometric Functions"]

QUALITY_WORD = {
    "verified visual transcription": "verified (transcribed from the page image)",
    "dual-pass OCR extraction, not visually QA'd": "ocr-extract (machine OCR, dual-pass, not visually checked)",
    "raw OCR reconstructed from the same page": "ocr-draft (raw OCR, not checked)",
    "no question text on this page": "marker only",
}


def add_field(doc, label, text, bold_text=False, color=None, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(label)
    r.bold = True
    if color:
        r.font.color.rgb = color
    r2 = p.add_run(text)
    r2.bold = bold_text
    return p


def hr(doc):
    """Thin horizontal rule under a paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFBFBF")
    borders.append(bottom)
    pPr.append(borders)


def main():
    recs = json.load(open(SRC))
    questions = [r for r in recs if r["prompt"]]
    markers = [r for r in recs if not r["prompt"]]

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)          # A4
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)
    for name, size in (("Heading 1", 20), ("Heading 2", 14), ("Heading 3", 12)):
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    # ---------------- cover ----------------
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(120)
    run = t.add_run("DIBY — Do It By Yourself")
    run.bold = True
    run.font.size = Pt(30)
    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Consolidated question bank — every DIBY item found in the OCR repository")
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    sub2 = doc.add_paragraph()
    sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub2.add_run(f"{len(questions)} questions  ·  {len(markers)} section markers  ·  "
                     f"{len(CHAPTER_ORDER)} chapters  ·  ARJUNA JEE 2027 class notes")
    r.font.size = Pt(11)

    doc.add_paragraph()
    stats = doc.add_paragraph()
    stats.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q = Counter(r["prompt_source"].split("(")[0].strip() for r in questions)
    n_verified = q.get("verified visual transcription", 0)
    n_extract = q.get("dual-pass OCR extraction, not visually QA'd", 0)
    n_draft = q.get("raw OCR reconstructed from the same page", 0)
    stats.add_run("Question text quality:  ").bold = True
    stats.add_run(f"{n_verified} verified   ·   {n_extract} ocr-extract   ·   {n_draft} ocr-draft")

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(24)
    rr = note.add_run(
        "Each entry gives the question, the options, the answer shown on the slide, the DIBY tag, "
        "and exactly where it comes from (chapter, lecture, source PDF and PDF page).\n"
        "Text was taken from the class-note slides; nothing was corrected or re-worded. "
        "Entries marked ocr-draft are raw OCR and must be checked against the source page."
    )
    rr.font.size = Pt(9.5)
    rr.italic = True
    rr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # ---------------- contents ----------------
    doc.add_page_break()
    doc.add_heading("Contents", level=1)
    for ch in CHAPTER_ORDER:
        ch_q = [r for r in questions if r["chapter"] == ch]
        if not ch_q:
            continue
        lects = sorted({r["lecture"] for r in ch_q})
        doc.add_paragraph(f"{ch} — {len(ch_q)} questions · lectures {lects[0]}–{lects[-1]}",
                          style="List Bullet")
    if markers:
        doc.add_paragraph(f"Appendix — {len(markers)} section markers with no question text "
                          "(divider slides)", style="List Bullet")

    # ---------------- body ----------------
    n = 0
    for ci, ch in enumerate(CHAPTER_ORDER):
        ch_q = sorted([r for r in questions if r["chapter"] == ch],
                      key=lambda r: (r["lecture"], r["pdf_page"], r["master_id"]))
        if not ch_q:
            continue
        doc.add_page_break() if ci else None
        doc.add_heading(ch, level=1)
        for lect in sorted({r["lecture"] for r in ch_q}):
            l_q = [r for r in ch_q if r["lecture"] == lect]
            src = l_q[0]["source_pdf"]
            doc.add_heading(f"Lecture {lect} — {src}", level=2)
            for r in l_q:
                n += 1
                tag = r["printed_label"] or (f"{r['internal_id']} (internal — not printed on the slide)"
                                             if r["internal_id"] else "no printed tag")
                h = doc.add_heading(f"Q{n} · {tag}", level=3)
                h.paragraph_format.space_before = Pt(10)
                h.paragraph_format.space_after = Pt(2)

                add_field(doc, "Question: ", r["prompt"])
                if r["options"]:
                    add_field(doc, "Options: ", r["options"])
                if r["answer"]:
                    add_field(doc, "Answer: ", r["answer"], bold_text=True)
                else:
                    add_field(doc, "Answer: ", "not shown on the slide",
                              color=RGBColor(0x88, 0x88, 0x88))
                tag_txt = r["printed_label"] or "— none printed —"
                if r["marker_text"] and r["marker_text"] != r["printed_label"]:
                    tag_txt += f"  |  marker as recorded: {r['marker_text']}"
                add_field(doc, "Tag: ", tag_txt)
                add_field(doc, "Where it comes from: ",
                          f"{r['source_pdf']} — {r['chapter']}, Lecture {r['lecture']}, "
                          f"PDF page {r['pdf_page']}  |  {QUALITY_WORD.get(r['prompt_source'].split('(')[0].strip(), r['prompt_source'])}")
                if r["repeat_of"]:
                    add_field(doc, "Repeat: ", f"the same question also appears at {r['repeat_of']}",
                              color=RGBColor(0x88, 0x55, 0x00))
                hr(doc)

    # ---------------- appendix ----------------
    if markers:
        doc.add_page_break()
        doc.add_heading("Appendix — section markers", level=1)
        p = doc.add_paragraph()
        p.add_run("These slides carry a DIBY / “Do It By Yourself” mark but no question text of their own — "
                  "they are divider or note slides. The questions themselves start on the page noted "
                  "in the last column where given.").italic = True
        tbl = doc.add_table(rows=1, cols=6)
        tbl.style = "Light Grid Accent 1"
        for i, htxt in enumerate(["Tag / marker", "Chapter", "Lecture", "PDF page", "Source PDF", "What was found"]):
            cell = tbl.rows[0].cells[i]
            cell.text = ""
            cell.paragraphs[0].add_run(htxt).bold = True
        for r in sorted(markers, key=lambda r: (CHAPTER_ORDER.index(r["chapter"])
                                                if r["chapter"] in CHAPTER_ORDER else 99,
                                                r["lecture"], r["pdf_page"])):
            cells = tbl.add_row().cells
            vals = [r["printed_label"] or r["internal_id"] or "—", r["chapter"], r["lecture"],
                    str(r["pdf_page"]), r["source_pdf"],
                    (r["uncertainty"] or "marker only")[:180]]
            for c, v in zip(cells, vals):
                c.text = ""
                c.paragraphs[0].add_run(v).font.size = Pt(8.5)

    doc.save(OUT)
    print("wrote", OUT, os.path.getsize(OUT), "bytes")
    print("questions:", len(questions), "| markers:", len(markers), "| numbered Q1..Q%d" % n)


if __name__ == "__main__":
    main()
